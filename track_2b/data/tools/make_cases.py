"""Generate the fictitious evaluation cases (PDF, DOCX, TXT, degraded scan) with their expected answers.

All companies and persons are invented. Every document carries a FICTITIOUS banner.
Run: python make_cases.py   (needs reportlab, python-docx, pillow; development tool, not used at runtime)
"""
from __future__ import annotations

import json
import random
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1] / "cases"
BANNER = {
    "de": "FIKTIVES TESTDOKUMENT. Alle Namen, Firmen und Zahlen sind frei erfunden.",
    "fr": "DOCUMENT DE TEST FICTIF. Tous les noms, sociétés et chiffres sont inventés.",
    "it": "DOCUMENTO DI PROVA FITTIZIO. Tutti i nomi, le società e le cifre sono inventati.",
}


def pdf(path: Path, lang: str, pages: list[list[str]]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for lines in pages:
        y = 800
        c.setFont("Helvetica-Oblique", 8)
        c.drawString(50, y, BANNER[lang])
        y -= 28
        for line in lines:
            if line.startswith("# "):
                c.setFont("Helvetica-Bold", 12)
                line = line[2:]
            else:
                c.setFont("Helvetica", 10)
            c.drawString(50, y, line)
            y -= 17
        c.showPage()
    c.save()


def txt(path: Path, lang: str, pages: list[list[str]]) -> None:
    out = [BANNER[lang]]
    for i, lines in enumerate(pages, 1):
        marker = {"de": "Seite", "fr": "Page", "it": "Pagina"}[lang]
        out += [f"{marker} {i}", *[ln.removeprefix("# ") for ln in lines]]
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def docx_file(path: Path, lang: str, title: str, intro: list[str], table: list[list[str]], after: list[str]) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph(BANNER[lang])
    d.add_heading(title, level=1)
    for p in intro:
        d.add_paragraph(p)
    t = d.add_table(rows=0, cols=len(table[0]))
    for row in table:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    for p in after:
        d.add_paragraph(p)
    d.save(str(path))


def scan(path: Path, lang: str, lines: list[str], seed: int = 7) -> None:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    rnd = random.Random(seed)
    font_path = next((p for p in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
                      if Path(p).exists()), None)
    font = ImageFont.truetype(font_path, 30) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 22) if font_path else font
    img = Image.new("L", (1700, 1100), 238)
    draw = ImageDraw.Draw(img)
    draw.text((80, 60), BANNER[lang], fill=60, font=small)
    y = 130
    for line in lines:
        draw.text((80 + rnd.randint(-3, 3), y), line.removeprefix("# "), fill=rnd.randint(20, 70), font=font)
        y += 52
    for _ in range(2500):  # speckle noise
        x, yy = rnd.randint(0, 1699), rnd.randint(0, 1099)
        draw.point((x, yy), fill=rnd.randint(90, 200))
    img = img.rotate(1.2, fillcolor=238).filter(ImageFilter.GaussianBlur(1.1))
    img.save(path, quality=55)


CASES: list[dict] = []


def case(cid: str, target: str, lang: str, docs: list[tuple], expected: dict) -> None:
    CASES.append({"id": cid, "target": target, "lang": lang, "docs": docs, "expected": expected})


def bo(name, basis, band, status="settled", **kw):
    return {"name": name, "basis": basis, "band": band, "status": status, **kw}


def h(holder, company, units=None, percent=None, held_for=None):
    return {"holder": holder, "company": company, "units_held": units, "percent": percent, "held_for": held_for}


# 1. Direct holdings, German, text PDF
case("case01_direct_de", "Rigiblick Immobilien AG", "de",
     [("pdf", "aktienbuch.pdf", [[
         "# Aktienbuch der Rigiblick Immobilien AG, Arth",
         "Aktienkapital CHF 200'000, eingeteilt in 2'000 Namenaktien zu je CHF 100.",
         "Stand per 30. September 2026",
         "Aktionär                          Wohnsitz              Anzahl Aktien",
         "Reto Fankhauser, geb. 14.02.1968, Schweizer   3000 Bern    1'100",
         "Claudia Wyss, geb. 03.07.1975, Schweizerin    6300 Zug       600",
         "Simon Gerber, geb. 22.11.1990, Schweizer      6410 Goldau    300",
     ]])],
     {"beneficial_owners": [bo("Reto Fankhauser", "direct", "50_75"), bo("Claudia Wyss", "direct", "25_50")],
      "review_reasons": [],
      "holdings": [h("Reto Fankhauser", "Rigiblick Immobilien AG", 1100), h("Claudia Wyss", "Rigiblick Immobilien AG", 600),
                   h("Simon Gerber", "Rigiblick Immobilien AG", 300)],
      "notes": "Two direct holders above 25 %, one below."})

# 2. Indirect through a holding, French, two PDFs
case("case02_indirect_fr", "Boulangerie des Alpes SA", "fr",
     [("pdf", "01_registre_boulangerie.pdf", [[
         "# Registre des actions de Boulangerie des Alpes SA, Bulle",
         "Capital-actions de CHF 100'000, divisé en 1'000 actions nominatives de CHF 100.",
         "Actionnaire                                   Domicile           Nombre d'actions",
         "Alpes Gourmet Holding SA                      1630 Bulle                   700",
         "Julien Perrin, né le 9 mai 1979, suisse       1700 Fribourg                300",
     ]]),
      ("pdf", "02_registre_holding.pdf", [[
          "# Registre des actions de Alpes Gourmet Holding SA, Bulle",
          "Capital-actions de CHF 50'000, divisé en 500 actions nominatives.",
          "Valérie Chappuis, née le 2 mars 1972, suisse, 1630 Bulle : 260 actions",
          "Nicolas Favre, né le 18 octobre 1981, suisse, 1700 Fribourg : 240 actions",
      ]])],
     {"beneficial_owners": [bo("Julien Perrin", "direct", "25_50"), bo("Valérie Chappuis", "indirect", "50_75")],
      "review_reasons": [],
      "holdings": [h("Alpes Gourmet Holding SA", "Boulangerie des Alpes SA", 700),
                   h("Julien Perrin", "Boulangerie des Alpes SA", 300),
                   h("Valérie Chappuis", "Alpes Gourmet Holding SA", 260),
                   h("Nicolas Favre", "Alpes Gourmet Holding SA", 240)],
      "notes": "Indirect control: 52 % of a holding that owns 70 %. Band from the holding's 70 %."})

# 3. Two-level chain, German, three documents (PDF, TXT, PDF)
case("case03_two_level_de", "Engelberg Bergbahnen AG", "de",
     [("pdf", "01_aktienbuch_bergbahnen.pdf", [[
         "# Aktienbuch der Engelberg Bergbahnen AG",
         "Das Aktienkapital ist eingeteilt in 10'000 Namenaktien.",
         "Titlis Invest AG, Engelberg                     8'000 Aktien",
         "Peter Amstutz, geb. 01.04.1960, Schweizer, Stans  2'000 Aktien",
     ]]),
      ("txt", "02_aktienbuch_titlis.txt", [[
          "# Aktienbuch der Titlis Invest AG",
          "Aktienkapital CHF 100'000, eingeteilt in 1'000 Namenaktien.",
          "Alpina Familienholding GmbH, Engelberg | 650",
          "Karin Odermatt, geb. 12.12.1985, Schweizerin, Luzern | 350",
      ]]),
      ("pdf", "03_gesellschafter_alpina.pdf", [[
          "# Gesellschafterliste der Alpina Familienholding GmbH",
          "Stammkapital CHF 20'000, eingeteilt in 200 Stammanteile.",
          "Martin Odermatt, geb. 05.06.1958, Schweizer, Engelberg: 200 Stammanteile",
      ]])],
     {"beneficial_owners": [bo("Martin Odermatt", "indirect", "75_100", chain_info_required=True)],
      "review_reasons": [],
      "holdings": [h("Titlis Invest AG", "Engelberg Bergbahnen AG", 8000), h("Peter Amstutz", "Engelberg Bergbahnen AG", 2000),
                   h("Alpina Familienholding GmbH", "Titlis Invest AG", 650), h("Karin Odermatt", "Titlis Invest AG", 350),
                   h("Martin Odermatt", "Alpina Familienholding GmbH", 200)],
      "notes": "Chain of two intermediaries; TJPV art. 15 chain information required."})

# 4. Voting pact, German
case("case04_pact_de", "Zürichsee Werft AG", "de",
     [("pdf", "01_aktienbuch.pdf", [[
         "# Aktienbuch der Zürichsee Werft AG, Wädenswil",
         "Aktienkapital eingeteilt in 500 Namenaktien.",
         "Beat Kälin, geb. 30.08.1966, Schweizer, Wädenswil       325 Aktien",
         "Andrea Brunner, geb. 17.01.1979, Schweizerin, Horgen    100 Aktien",
         "Felix Brunner, geb. 09.09.1982, Schweizer, Thalwil       75 Aktien",
     ]]),
      ("pdf", "02_aktionaerbindungsvertrag.pdf", [[
          "# Aktionärbindungsvertrag vom 2. Mai 2025",
          "zwischen Andrea Brunner und Felix Brunner betreffend die Zürichsee Werft AG",
          "Art. 3 Stimmbindung",
          "Felix Brunner übt seine Stimmrechte an der Generalversammlung der",
          "Zürichsee Werft AG stets nach den Weisungen von Andrea Brunner aus.",
      ]])],
     {"beneficial_owners": [bo("Beat Kälin", "direct", "50_75"),
                            bo("Andrea Brunner", "other_means", "25_50", "pending_review")],
      "review_reasons": ["concert"],
      "holdings": [h("Beat Kälin", "Zürichsee Werft AG", 325), h("Andrea Brunner", "Zürichsee Werft AG", 100),
                   h("Felix Brunner", "Zürichsee Werft AG", 75)],
      "notes": "20 % plus 15 % bound by a voting agreement: candidate pending review (TJPV art. 3, 4)."})

# 5. Fiduciary holding, French, TXT
case("case05_fiduciary_fr", "Cave du Lavaux Sàrl", "fr",
     [("txt", "registre_parts.txt", [[
         "# Registre des parts sociales de Cave du Lavaux Sàrl, Cully",
         "Capital social de CHF 20'000, divisé en 100 parts sociales.",
         "Laurent Rey, né le 11 novembre 1970, suisse, 1096 Cully : 40 parts.",
         "Fiduciaire Rhône SA, Sion : 60 parts, détenues à titre fiduciaire pour le compte de Sophie Bonvin.",
         "Sophie Bonvin, née le 25 janvier 1983, suisse, 1950 Sion.",
     ]])],
     {"beneficial_owners": [bo("Laurent Rey", "direct", "25_50"),
                            bo("Sophie Bonvin", "direct", "50_75", "pending_review")],
      "review_reasons": ["fiduciary"],
      "holdings": [h("Laurent Rey", "Cave du Lavaux Sàrl", 40),
                   h("Fiduciaire Rhône SA", "Cave du Lavaux Sàrl", 60, held_for="Sophie Bonvin")],
      "notes": "Nominee never proposed; principal pending review (TJPV art. 3 para. 2 let. e)."})

# 6. Person just below the threshold, Italian, PDF
case("case06_below_it", "Pasticceria Ceresio SA", "it",
     [("pdf", "libro_azioni.pdf", [[
         "# Libro delle azioni della Pasticceria Ceresio SA, Lugano",
         "Capitale azionario di CHF 100'000, suddiviso in 1'000 azioni nominative.",
         "Marco Bernasconi, nato il 3 marzo 1969, svizzero, 6900 Lugano        400 azioni",
         "Giulia Rossi, nata il 14 luglio 1984, italiana, 6830 Chiasso          240 azioni",
         "Luca Ferrari, nato il 2 febbraio 1977, svizzero, 6500 Bellinzona     200 azioni",
         "Anna Galli, nata il 30 settembre 1991, svizzera, 6600 Locarno        160 azioni",
     ]])],
     {"beneficial_owners": [bo("Marco Bernasconi", "direct", "25_50")],
      "review_reasons": [],
      "holdings": [h("Marco Bernasconi", "Pasticceria Ceresio SA", 400), h("Giulia Rossi", "Pasticceria Ceresio SA", 240),
                   h("Luca Ferrari", "Pasticceria Ceresio SA", 200), h("Anna Galli", "Pasticceria Ceresio SA", 160)],
      "notes": "Giulia Rossi holds 24 %: not a beneficial owner."})

# 7. Nobody above 25 %, default manager, German
case("case07_nobody_de", "Software Kollektiv AG", "de",
     [("pdf", "01_aktienbuch.pdf", [[
         "# Aktienbuch der Software Kollektiv AG, Winterthur",
         "Aktienkapital eingeteilt in 1'000 Namenaktien.",
         "Lea Bachmann, Winterthur        200 Aktien",
         "Jonas Keller, Zürich            200 Aktien",
         "Mira Steiner, Uster             200 Aktien",
         "David Huber, Wil                200 Aktien",
         "Sara Frei, Frauenfeld           200 Aktien",
     ]]),
      ("pdf", "02_handelsregisterauszug.pdf", [[
          "# Auszug Verwaltungsrat der Software Kollektiv AG",
          "Nadine Schmid, geb. 21.06.1974, Schweizerin, in Winterthur, Präsidentin des Verwaltungsrates",
          "Jonas Keller, geb. 08.03.1988, Schweizer, in Zürich, Mitglied des Verwaltungsrates",
      ]])],
     {"beneficial_owners": [bo("Nadine Schmid", "default_manager", None, "pending_review")],
      "review_reasons": ["default_manager"],
      "holdings": [h(n, "Software Kollektiv AG", 200) for n in
                   ("Lea Bachmann", "Jonas Keller", "Mira Steiner", "David Huber", "Sara Frei")],
      "notes": "Five holders at 20 %: TJPG art. 4 para. 2, chair of the board."})

# 8. Multiplication trap, French
case("case08_trap_fr", "Transports Jura SA", "fr",
     [("pdf", "01_registre_transports.pdf", [[
         "# Registre des actions de Transports Jura SA, Delémont",
         "Le capital-actions est divisé en 1'000 actions nominatives.",
         "Jura Logistique SA, Delémont : 1'000 actions",
         "Président du conseil d'administration : Pierre Monnier, né le 6 août 1962, suisse, Delémont",
     ]]),
      ("pdf", "02_registre_logistique.pdf", [[
          "# Registre des actions de Jura Logistique SA, Delémont",
          "Capital-actions divisé en 100 actions nominatives.",
          "Pierre Monnier, Delémont : 40 actions",
          "Hélène Monnier, Porrentruy : 30 actions",
          "Paul Monnier, Moutier : 30 actions",
      ]])],
     {"beneficial_owners": [bo("Pierre Monnier", "default_manager", None, "pending_review")],
      "review_reasons": ["default_manager"],
      "holdings": [h("Jura Logistique SA", "Transports Jura SA", 1000), h("Pierre Monnier", "Jura Logistique SA", 40),
                   h("Hélène Monnier", "Jura Logistique SA", 30), h("Paul Monnier", "Jura Logistique SA", 30)],
      "notes": "40 % of a 100 % holding: product 40 % but no control of the intermediate (TJPV art. 2)."})

# 9. Inconsistent figures, German
case("case09_inconsistent_de", "Holzbau Emmental AG", "de",
     [("pdf", "aktienbuch.pdf", [[
         "# Aktienbuch der Holzbau Emmental AG, Langnau",
         "Aktienkapital eingeteilt in 1'000 Namenaktien.",
         "Urs Lehmann, geb. 19.10.1963, Schweizer, Langnau     600 Aktien",
         "Erika Lehmann, geb. 07.05.1966, Schweizerin, Langnau 300 Aktien",
     ]])],
     {"beneficial_owners": [bo("Urs Lehmann", "direct", "50_75", "pending_review"),
                            bo("Erika Lehmann", "direct", "25_50", "pending_review")],
      "review_reasons": ["inconsistent_figures"],
      "holdings": [h("Urs Lehmann", "Holzbau Emmental AG", 600), h("Erika Lehmann", "Holzbau Emmental AG", 300)],
      "notes": "Holdings add up to 900 of 1'000 shares."})

# 10. Undocumented foreign holder, French
case("case10_foreign_fr", "Montres du Léman SA", "fr",
     [("pdf", "registre.pdf", [[
         "# Registre des actions de Montres du Léman SA, Genève",
         "Capital-actions divisé en 1'000 actions nominatives.",
         "Pacific Horizon Holdings Ltd, Singapour : 700 actions",
         "Claire Dubois, née le 13 avril 1980, française, 1204 Genève : 300 actions",
     ]])],
     {"beneficial_owners": [bo("Claire Dubois", "direct", "25_50")],
      "review_reasons": ["foreign_or_undocumented"],
      "missing": [{"what": "ownership_proof", "about": "Pacific Horizon Holdings Ltd"}],
      "holdings": [h("Pacific Horizon Holdings Ltd", "Montres du Léman SA", 700), h("Claire Dubois", "Montres du Léman SA", 300)],
      "notes": "Owners of the foreign holding unknown: request documents."})

# 11. DOCX, Italian
case("case11_docx_it", "Costruzioni Ticino SA", "it",
     [("docx", "libro_azioni.docx", ("Libro delle azioni della Costruzioni Ticino SA",
                                     ["Capitale azionario suddiviso in 1'000 azioni nominative."],
                                     [["Azionista", "Domicilio", "Azioni"],
                                      ["Paolo Conti", "6500 Bellinzona", "550"],
                                      ["Sara Conti", "6512 Giubiasco", "450"]],
                                     ["Presidente del consiglio di amministrazione: Paolo Conti."]))],
     {"beneficial_owners": [bo("Paolo Conti", "direct", "50_75"), bo("Sara Conti", "direct", "25_50")],
      "review_reasons": [],
      "holdings": [h("Paolo Conti", "Costruzioni Ticino SA", 550), h("Sara Conti", "Costruzioni Ticino SA", 450)],
      "notes": "Share register as a DOCX table."})

# 12. Poor scan, German
case("case12_scan_de", "Garage Aaretal AG", "de",
     [("scan", "aktienbuch_scan.jpg", [
         "# Aktienbuch der Garage Aaretal AG, Thun",
         "Aktienkapital eingeteilt in 100 Namenaktien.",
         "Thomas Wenger, Thun          70 Aktien",
         "Monika Wenger, Steffisburg   30 Aktien",
     ])],
     {"beneficial_owners": [bo("Thomas Wenger", "direct", "50_75", "any"), bo("Monika Wenger", "direct", "25_50", "any")],
      "review_reasons": [],
      "holdings": [h("Thomas Wenger", "Garage Aaretal AG", 70), h("Monika Wenger", "Garage Aaretal AG", 30)],
      "notes": "Degraded scan (blur, rotation, noise); status may be pending if OCR confidence is low."})


def demo_cases() -> None:
    case("case13_chain_pact_de", "Alpenblick Holding AG", "de",
         [("txt", "aktienbuch.txt", [[
             "# AKTIENBUCH der Alpenblick Holding AG, mit Sitz in Luzern",
             "Stand: 30. September 2026",
             "Aktienkapital: CHF 100'000, eingeteilt in 1'000 Namenaktien zu nominal CHF 100.",
             "Nr. | Aktionär | Adresse | Anzahl Aktien | Nummern",
             "1 | Petra Imhof, geb. 12.03.1971, Schweizer Bürgerin | Seestrasse 4, 6006 Luzern | 520 | 1 bis 520",
             "2 | Bergkristall Beteiligungen GmbH | Bahnhofplatz 2, 6300 Zug | 300 | 521 bis 820",
             "3 | Lukas Ammann, geb. 02.11.1985, Schweizer Bürger | Dorfweg 9, 6010 Kriens | 180 | 821 bis 1000",
         ], [
             "Auszug aus dem Aktionärbindungsvertrag vom 15. Januar 2025",
             "Art. 7 Stimmbindung: Lukas Ammann verpflichtet sich, seine Stimmrechte an allen",
             "Generalversammlungen gemäss den Weisungen von Petra Imhof auszuüben.",
             "Die Gesellschafterin der Bergkristall Beteiligungen GmbH ist zu 100 % Frau Sandra Keller,",
             "wohnhaft in Zug.",
         ]])],
         {"beneficial_owners": [bo("Petra Imhof", "direct", "50_75"), bo("Sandra Keller", "indirect", "25_50")],
          "review_reasons": ["concert"],
          "holdings": [h("Petra Imhof", "Alpenblick Holding AG", 520),
                       h("Bergkristall Beteiligungen GmbH", "Alpenblick Holding AG", 300),
                       h("Lukas Ammann", "Alpenblick Holding AG", 180),
                       h("Sandra Keller", "Bergkristall Beteiligungen GmbH", percent=100)],
          "notes": "Table plus prose: 100 % written as a percentage; voting agreement on a settled owner."})
    case("case14_chain_fiduciary_fr", "Horlogerie du Jorat SA", "fr",
         [("txt", "registre.txt", [[
             "# REGISTRE DES ACTIONS de Horlogerie du Jorat SA, siège à Lausanne",
             "Capital-actions : CHF 250'000, divisé en 2'500 actions nominatives de CHF 100.",
             "Actionnaire | Domicile | Nombre d'actions",
             "Jorat Participations Sàrl | Rue du Lac 11, 1003 Lausanne | 1'500",
             "Marc Dufour, né le 4 juin 1968, de nationalité suisse | Chemin des Vignes 3, 1095 Lutry | 600",
             "Élise Moret, née le 21 août 1990, de nationalité française | Avenue de la Gare 8, 1800 Vevey | 400",
         ], [
             "Extrait du registre des parts sociales de Jorat Participations Sàrl :",
             "capital social de CHF 20'000, 200 parts sociales.",
             "Anne Rochat, Lausanne : 120 parts. Fiduciaire Léman SA, à titre fiduciaire pour le compte de",
             "Thomas Berger, Genève : 80 parts (convention de fiducie du 3 mars 2024).",
             "Président du conseil d'administration de Horlogerie du Jorat SA : Marc Dufour.",
         ]])],
         {"beneficial_owners": [bo("Anne Rochat", "indirect", "50_75")],
          "review_reasons": ["fiduciary"],
          "holdings": [h("Jorat Participations Sàrl", "Horlogerie du Jorat SA", 1500),
                       h("Marc Dufour", "Horlogerie du Jorat SA", 600), h("Élise Moret", "Horlogerie du Jorat SA", 400),
                       h("Anne Rochat", "Jorat Participations Sàrl", 120),
                       h("Fiduciaire Léman SA", "Jorat Participations Sàrl", 80, held_for="Thomas Berger")],
          "notes": "Wrapped sentence; fiduciary holding without control; Dufour at 24 % is not an owner."})


def main() -> None:
    demo_cases()
    for c in CASES:
        d = ROOT / c["id"]
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
        for kind, name, content in c["docs"]:
            if kind == "pdf":
                pdf(d / name, c["lang"], content)
            elif kind == "txt":
                txt(d / name, c["lang"], content)
            elif kind == "docx":
                docx_file(d / name, c["lang"], *content)
            elif kind == "scan":
                scan(d / name, c["lang"], content)
        (d / "case.json").write_text(json.dumps({"target": c["target"], "fictitious": True, "language": c["lang"]},
                                                ensure_ascii=False, indent=1), encoding="utf-8")
        exp = {"fictitious": True, "licence": "CC-BY-4.0", "target": c["target"], "language": c["lang"],
               **c["expected"]}
        (d / "expected.json").write_text(json.dumps(exp, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(CASES)} cases written to {ROOT}")


if __name__ == "__main__":
    main()
